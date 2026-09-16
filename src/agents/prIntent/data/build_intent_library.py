#!/usr/bin/env python3
"""
Builds the PR Intent knowledge base consumed by the multi-agent PR Intent
feature (src/agents/prIntent/**) from the three source spreadsheets:

  - PR Intent Chatbot Question.xlsx        (operational schema: trigger
    patterns, API query, calculation, answer template, chart, evidence,
    follow-up, action)
  - PR Intent Library Details.xlsx          (richer schema: persona,
    question variations, data sources, expected answer, visualization,
    recommended action)
  - PR Solutions Chatbot Question Library.xlsx (a flat library of ~300
    example natural-language questions grouped under KPI pillars/personas
    - used only to widen the trigger-phrase corpus for the classifier)

Both intent sheets share the same ID scheme (MM-xxx, PI-xxx, NI-xxx,
RI-xxx, CI-xxx, CA-xxx, ST-xxx, PE-xxx, RK-xxx, EX-xxx, CX-xxx) so records
are merged by ID. The merged result is grouped into 7 "agent domains" -
one per specialised PR-analytics agent - mirroring how the source sheets
themselves are organised, per the master architecture described in
"PR Intent Chatbot Question.xlsx" (Intent Detection -> KPI Library match
-> Intent JSON -> data layer -> KPI calc -> Insight agent -> evidence +
chart -> answer -> follow-up).

Re-run this script (`python3 build_intent_library.py`) whenever the source
xlsx files change. Requires `pip install openpyxl`. Output:
  - intentLibrary.json      merged, structured intent records by domain
  - rawQuestionCorpus.json  extra natural-language trigger phrases by domain
"""
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl --break-system-packages")

UPLOADS = Path("/sessions/epic-kind-babbage/mnt/uploads")
OUT_DIR = Path(__file__).resolve().parent

FILE1 = UPLOADS / "PR Intent Chatbot Question.xlsx"
FILE2 = UPLOADS / "PR Intent Library Details.xlsx"
FILE3 = UPLOADS / "PR Solutions Chatbot Question Library.xlsx"

# Every sheet in file1/file2 maps to one of 7 segregated agent domains.
# Story Intelligence + People & Spokesperson fold into Narrative
# Intelligence (they're both about "what we said / who said it and how it
# travelled"); Reputation Risk folds into Reputation Index; Executive
# Strategic + Cross-KPI fold into a 7th "executive cross-kpi" agent whose
# whole job is to fan out across the other 6 and merge their answers.
SHEET_TO_DOMAIN = {
    "Media Measurement": "mediaMeasurement",
    "Media Measurements": "mediaMeasurement",
    "PR Impact": "prImpact",
    "Narrative Intelligence": "narrativeIntelligence",
    "Reputation Index": "reputationIndex",
    "Competitive PR": "competitiveIntelligence",
    "Competitive Intelligence": "competitiveIntelligence",
    "CAMPAIGN INTELLIGENCE": "campaign",
    "Campaign Intents": "campaign",
    "Story Intelligence": "narrativeIntelligence",
    "People & Spokesperson": "narrativeIntelligence",
    "Reputation Risk": "reputationIndex",
    "Executive & Strategic Question": "executiveCrossKpi",
    "Executive Startegic": "executiveCrossKpi",  # sic - typo in source workbook
    "Cross-KPI Question": "executiveCrossKpi",
    "Cross-KPI Intents": "executiveCrossKpi",
}

DOMAIN_LABELS = {
    "mediaMeasurement": "Media Measurement",
    "prImpact": "PR Impact",
    "narrativeIntelligence": "Narrative Intelligence",
    "reputationIndex": "Reputation Index",
    "competitiveIntelligence": "Competitive Intelligence",
    "campaign": "Campaign Intelligence",
    "executiveCrossKpi": "Executive & Cross-KPI",
}

# "PR Solutions Chatbot Question Library.xlsx" groups its ~300 example
# questions under free-text pillar/persona headings rather than IDs.
MAJOR_TO_DOMAIN = {
    "1. MEDIA MEASUREMENT": "mediaMeasurement",
    "2. PR IMPACT": "prImpact",
    "3. NARRATIVE INTELLIGENCE / MESSAGE CONGRUENCE": "narrativeIntelligence",
    "4. REPUTATION INDEX": "reputationIndex",
    "5. COMPETITIVE PR QUESTIONS": "competitiveIntelligence",
    "6. CAMPAIGN PERFORMANCE QUESTIONS": "campaign",
    "7. EXECUTIVE / CCO QUESTIONS": "executiveCrossKpi",
    "8. PR AGENCY QUESTIONS": "executiveCrossKpi",
    "9. JUNIOR PR EXECUTIVE QUESTIONS": "mediaMeasurement",
    "10. PR EXECUTIVE / MANAGER QUESTIONS": "executiveCrossKpi",
}


def split_list(value, seps=(";", "/", ",")):
    """Split a free-text spreadsheet cell into a clean list of phrases."""
    if value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    pattern = "|".join(re.escape(s) for s in seps)
    parts = re.split(pattern, text)
    return [p.strip().strip('"').strip() for p in parts if p.strip()]


def extract_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    for i, r in enumerate(rows):
        if r and any(isinstance(c, str) and c.strip() == "ID" for c in r):
            header_idx = i
            break
    if header_idx is None:
        return []
    header = [(c.strip() if isinstance(c, str) else c) for c in rows[header_idx]]
    out = []
    for r in rows[header_idx + 1:]:
        if not r or not r[0]:
            continue
        rec = {h: v for h, v in zip(header, r) if h is not None}
        out.append(rec)
    return out


def build_from_file1():
    """intent_id -> operational-schema fields."""
    wb = openpyxl.load_workbook(FILE1, data_only=True)
    merged = {}
    for sheet in wb.sheetnames:
        domain = SHEET_TO_DOMAIN.get(sheet.strip())
        if not domain:
            continue
        for rec in extract_sheet(wb[sheet]):
            intent_id = str(rec.get("ID", "")).strip()
            if not intent_id:
                continue
            merged[intent_id] = {
                "id": intent_id,
                "domain": domain,
                "name": rec.get("Intent Name"),
                "kpi": rec.get("KPI"),
                "triggerPatterns": split_list(rec.get("Trigger Patterns")),
                "apiQuery": rec.get("API/Data Query"),
                "filters": split_list(rec.get("Filters")),
                "calculation": rec.get("Calculation"),
                "answerTemplate": rec.get("Answer Template"),
                "chartHint": rec.get("Chart"),
                "evidence": rec.get("Evidence"),
                "followUp": split_list(rec.get("Follow-up"), seps=(";",)),
                "action": split_list(rec.get("Action"), seps=(";",)),
            }
    return merged


def build_from_file2():
    """intent_id -> richer library-details fields (persona, question
    variations, data sources, visualization, recommended action)."""
    wb = openpyxl.load_workbook(FILE2, data_only=True)
    merged = {}
    for sheet in wb.sheetnames:
        domain = SHEET_TO_DOMAIN.get(sheet.strip())
        if not domain:
            continue
        for rec in extract_sheet(wb[sheet]):
            intent_id = str(rec.get("ID", "")).strip()
            if not intent_id:
                continue
            if "Question Variations" in rec:
                merged[intent_id] = {
                    "id": intent_id,
                    "domain": domain,
                    "name": rec.get("Intent"),
                    "kpi": rec.get("KPI"),
                    "subKpi": rec.get("Sub-KPI"),
                    "personas": split_list(rec.get("User Persona")),
                    "questionVariations": split_list(rec.get("Question Variations")),
                    "requiredFilters": split_list(rec.get("Required Filters")),
                    "dataSources": split_list(rec.get("Data Sources")),
                    "calculation": rec.get("Calculation"),
                    "expectedAnswer": rec.get("Expected Answer"),
                    "visualization": rec.get("Visualization"),
                    "followUpQuestions": split_list(rec.get("Follow-up Questions"), seps=(";",)),
                    "recommendedAction": rec.get("Recommended Action"),
                }
            else:
                # Cross-KPI Intents sheet has a different, narrower schema.
                merged[intent_id] = {
                    "id": intent_id,
                    "domain": domain,
                    "name": rec.get("Intent"),
                    "kpiCombination": split_list(rec.get("KPI Combination"), seps=(";", "+", ",")),
                    "questionVariations": split_list(rec.get("User Question"), seps=(";",)),
                    "calculation": rec.get("Calculation / Reasoning"),
                    "expectedAnswer": rec.get("Expected Answer"),
                }
    return merged


def build_raw_corpus():
    """Extra example questions (no IDs) grouped by domain, used only to
    widen the classifier's trigger-phrase corpus."""
    wb = openpyxl.load_workbook(FILE3, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))

    major_re = re.compile(r"^\d+\.\s+[A-Z].*")
    minor_re = re.compile(r"^[A-Z]\.\s+.*")

    current_major = None
    current_minor = None
    by_domain = {}
    for r in rows:
        c_val = r[2] if len(r) > 2 else None
        d_val = r[3] if len(r) > 3 else None
        if isinstance(c_val, str) and major_re.match(c_val.strip()):
            current_major = c_val.strip()
            continue
        if isinstance(c_val, str) and minor_re.match(c_val.strip()) and len(c_val.strip()) < 80:
            current_minor = c_val.strip()
            continue
        if isinstance(c_val, (int, float)) and isinstance(d_val, str) and d_val.strip():
            domain = MAJOR_TO_DOMAIN.get(current_major)
            if not domain:
                continue
            by_domain.setdefault(domain, []).append({
                "question": d_val.strip(),
                "section": current_minor,
            })
    return by_domain


def merge():
    f1 = build_from_file1()
    f2 = build_from_file2()
    all_ids = sorted(set(f1) | set(f2))

    domains = {key: {"label": DOMAIN_LABELS[key], "intents": []} for key in DOMAIN_LABELS}

    for intent_id in all_ids:
        a = f1.get(intent_id, {})
        b = f2.get(intent_id, {})
        domain = a.get("domain") or b.get("domain")
        if domain not in domains:
            continue
        merged_intent = {
            "id": intent_id,
            "domain": domain,
            "name": a.get("name") or b.get("name"),
            "kpi": a.get("kpi") or b.get("kpi"),
            "subKpi": b.get("subKpi"),
            "personas": b.get("personas", []),
            "triggerPatterns": sorted(set(
                [p.lower() for p in a.get("triggerPatterns", [])]
                + [p.lower() for p in b.get("questionVariations", [])]
            )),
            "requiredFilters": sorted(set(a.get("filters", []) + b.get("requiredFilters", []))),
            "dataSources": b.get("dataSources", []),
            "apiQuery": a.get("apiQuery"),
            "calculation": a.get("calculation") or b.get("calculation"),
            "answerTemplate": a.get("answerTemplate") or b.get("expectedAnswer"),
            "chartHint": a.get("chartHint") or b.get("visualization"),
            "evidence": a.get("evidence"),
            "followUp": sorted(set(a.get("followUp", []) + b.get("followUpQuestions", []))),
            "action": a.get("action") or ([b["recommendedAction"]] if b.get("recommendedAction") else []),
            "kpiCombination": b.get("kpiCombination", []),
        }
        domains[domain]["intents"].append(merged_intent)

    return domains


def main():
    library = merge()
    corpus = build_raw_corpus()

    (OUT_DIR / "intentLibrary.json").write_text(json.dumps(library, indent=2, default=str))
    (OUT_DIR / "rawQuestionCorpus.json").write_text(json.dumps(corpus, indent=2, default=str))

    total = sum(len(d["intents"]) for d in library.values())
    print(f"Wrote {total} merged intents across {len(library)} domains -> intentLibrary.json")
    for key, d in library.items():
        print(f"  {key:24s} {len(d['intents']):3d} intents")
    print(f"Wrote raw question corpus -> rawQuestionCorpus.json "
          f"({sum(len(v) for v in corpus.values())} example questions)")


if __name__ == "__main__":
    main()
